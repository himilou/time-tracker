from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path
import tempfile
import unittest

from docx import Document

from timetracker.invoices import (
    create_invoice_docx,
    first_invoice_number,
    invoice_client_code,
    sessions_for_invoice,
)
from timetracker.models import WorkSession
from timetracker.storage import JsonStorage


class InvoiceTests(unittest.TestCase):
    def setUp(self) -> None:
        self.client_session = WorkSession(
            "Design",
            Decimal("60"),
            datetime(2026, 4, 2, 9, 0, tzinfo=timezone.utc),
            ended_at=datetime(2026, 4, 2, 10, 30, tzinfo=timezone.utc),
            client_name="Client & Co",
        )
        self.other_client_session = WorkSession(
            "Development",
            Decimal("80"),
            datetime(2026, 4, 2, 10, 0, tzinfo=timezone.utc),
            ended_at=datetime(2026, 4, 3, 11, 0, tzinfo=timezone.utc),
            client_name="Client B",
        )

    def test_filters_invoice_entries_by_client_and_end_date(self) -> None:
        selected = sessions_for_invoice(
            [self.client_session, self.other_client_session],
            "client & co",
            date(2026, 4, 2),
            date(2026, 4, 2),
        )

        self.assertEqual(selected, [self.client_session])

    def test_rejects_reversed_date_range(self) -> None:
        with self.assertRaises(ValueError):
            sessions_for_invoice(
                [self.client_session],
                "Client & Co",
                date(2026, 4, 30),
                date(2026, 4, 1),
            )

    def test_creates_editable_docx_invoice_for_selected_sessions(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "invoice.docx"
            second_session = WorkSession(
                "Support",
                Decimal("80"),
                datetime(2026, 4, 3, 9, 0, tzinfo=timezone.utc),
                ended_at=datetime(2026, 4, 3, 9, 30, tzinfo=timezone.utc),
                client_name="Client & Co",
            )
            total = create_invoice_docx(
                [self.client_session, second_session],
                "Client & Co",
                date(2026, 4, 1),
                date(2026, 4, 30),
                "2026-CC-001",
                output,
                invoice_date=date(2026, 5, 1),
            )

            self.assertEqual(total, Decimal("130.00"))
            document = Document(output)
            details = document.tables[0].cell(0, 1).text
            self.assertIn("2026-CC-001", details)
            self.assertIn("05/31/2026", details)
            self.assertIn("Status:  UNPAID", details)
            billing_heading = document.tables[0].cell(0, 0).paragraphs[0].runs[0]
            self.assertTrue(billing_heading.bold)
            self.assertEqual(str(billing_heading.font.color.rgb), "1F3A5F")
            self.assertIn(
                "Design - Work completed 2026-04-02",
                document.tables[1].cell(1, 0).text,
            )
            self.assertEqual(len(document.tables[1].rows), 3)
            self.assertIn(
                "Support - Work completed 2026-04-03",
                document.tables[1].cell(2, 0).text,
            )
            self.assertEqual(document.tables[2].cell(4, 1).text, "$130.00")
            self.assertEqual(document.tables[2].cell(1, 1).text, "$0.00")
            self.assertEqual(document.tables[2].cell(3, 1).text, "-$0.00")
            self.assertIn(
                "Payment due by May 31, 2026",
                "\n".join(paragraph.text for paragraph in document.paragraphs),
            )
            self.assertIn(
                "Please remit payment through PayPal to jameshardy1@gmail.com",
                "\n".join(
                    paragraph.text
                    for section in document.sections
                    for paragraph in section.footer.paragraphs
                ),
            )
            invoice_text = "\n".join(
                [paragraph.text for paragraph in document.paragraphs]
                + [
                    cell.text
                    for table in document.tables
                    for row in table.rows
                    for cell in row.cells
                ]
            )
            self.assertNotIn("Paid in full", invoice_text)
            self.assertNotIn("Lucky Liquor LLC", invoice_text)

    def test_invoice_client_code_and_template_sequence(self) -> None:
        self.assertEqual(invoice_client_code("Lucky Liquor LLC"), "LL")
        self.assertEqual(invoice_client_code("LuckyLiquor"), "LL")
        self.assertEqual(first_invoice_number(2026, "LL"), 222)
        self.assertEqual(first_invoice_number(2027, "LL"), 1)

    def test_invoice_number_sequences_are_persisted_by_year_and_client(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            storage = JsonStorage(Path(directory))
            self.assertEqual(
                storage.allocate_invoice_number(2026, "LL", 222), "2026-LL-222"
            )
            self.assertEqual(
                storage.allocate_invoice_number(2026, "LL", 222), "2026-LL-223"
            )
            self.assertEqual(
                storage.allocate_invoice_number(2026, "CC"), "2026-CC-001"
            )
            self.assertEqual(
                storage.allocate_invoice_number(2027, "LL"), "2027-LL-001"
            )

    def test_loads_invoice_sessions_across_month_files(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            storage = JsonStorage(Path(directory))
            session = WorkSession(
                "Design",
                Decimal("60"),
                datetime(2026, 4, 30, 23, 30, tzinfo=timezone.utc),
                ended_at=datetime(2026, 5, 1, 0, 30, tzinfo=timezone.utc),
                client_name="Client & Co",
            )
            storage.save_session(session)

            loaded_sessions = storage.load_all_sessions()

            self.assertEqual(len(loaded_sessions), 1)
            self.assertEqual(
                sessions_for_invoice(
                    loaded_sessions,
                    "Client & Co",
                    date(2026, 5, 1),
                    date(2026, 5, 1),
                ),
                [session],
            )


if __name__ == "__main__":
    unittest.main()
