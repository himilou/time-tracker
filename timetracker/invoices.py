"""Editable DOCX invoice generation from completed work sessions."""

from copy import deepcopy
from datetime import date, timedelta
from decimal import Decimal
from pathlib import Path
import re

from docx import Document
from docx.table import _Cell
from docx.text.paragraph import Paragraph

from .models import WorkSession

INVOICE_TEMPLATE = Path(__file__).parent / "templates" / "invoice_template.docx"


def sessions_for_invoice(
    sessions: list[WorkSession],
    client_name: str,
    start_date: date,
    end_date: date,
) -> list[WorkSession]:
    if end_date < start_date:
        raise ValueError("End date cannot be before start date.")
    return [
        session
        for session in sessions
        if session.is_complete
        and session.client_name.casefold() == client_name.casefold()
        and start_date <= session.ended_at.date() <= end_date
    ]


def invoice_client_code(client_name: str) -> str:
    separated_name = re.sub(r"(?<=[a-z0-9])(?=[A-Z])", " ", client_name)
    words = re.findall(r"[A-Za-z0-9]+", separated_name)
    words = [
        word for word in words if word.casefold() not in {"and", "the", "of", "llc"}
    ]
    if not words:
        raise ValueError("Client name must contain letters or numbers.")
    if len(words) == 1:
        return words[0][:2].upper()
    return "".join(word[0] for word in words).upper()


def first_invoice_number(year: int, client_code: str) -> int:
    template_details = Document(str(INVOICE_TEMPLATE)).tables[0].cell(0, 1).text
    match = re.search(r"Invoice No:\s*(\d{4})-([A-Z0-9]+)-(\d+)", template_details)
    if match and int(match.group(1)) == year and match.group(2) == client_code:
        return int(match.group(3)) + 1
    return 1


def create_invoice_docx(
    sessions: list[WorkSession],
    client_name: str,
    start_date: date,
    end_date: date,
    invoice_number: str,
    output_path: Path,
    invoice_date: date | None = None,
) -> Decimal:
    if not client_name.strip():
        raise ValueError("Select a client for the invoice.")
    selected_sessions = sessions_for_invoice(
        sessions, client_name, start_date, end_date
    )
    if not selected_sessions:
        raise ValueError("There are no completed sessions to invoice.")

    invoice_day = invoice_date or date.today()
    due_date = invoice_day + timedelta(days=30)
    ordered_sessions = sorted(
        selected_sessions, key=lambda item: (item.ended_at, item.project_name)
    )
    total = sum((session.total_pay for session in ordered_sessions), Decimal("0.00"))

    document = Document(str(INVOICE_TEMPLATE))
    details_table, items_table, totals_table = document.tables
    set_cell_lines(
        details_table.cell(0, 0),
        ["BILL TO", client_name, "", "", ""],
        bold_lines={0},
    )
    set_cell_lines(
        details_table.cell(0, 1),
        [
            f"Invoice No:  {invoice_number}",
            f"Invoice Date:  {invoice_day:%m/%d/%Y}",
            f"Due Date:  {due_date:%m/%d/%Y}",
            f"Service Period:  {start_date:%b %d, %Y} - {end_date:%b %d, %Y}",
            "Status:  UNPAID",
        ],
        bold_lines={4},
    )

    template_row = items_table.rows[1]._tr
    for row in list(items_table.rows[2:]):
        items_table._tbl.remove(row._tr)
    for index in range(1, len(ordered_sessions)):
        new_row = deepcopy(template_row)
        items_table._tbl.append(new_row)

    for row, session in zip(items_table.rows[1:], ordered_sessions):
        hours = (
            Decimal(session.billable_seconds // 60) / Decimal(60)
        ).quantize(Decimal("0.0001"))
        set_cell_lines(
            row.cells[0],
            [f"{session.project_name} - Work completed {session.ended_at:%Y-%m-%d}"],
        )
        set_cell_lines(row.cells[1], [f"{hours:.4f}"])
        set_cell_lines(row.cells[2], [format_currency(session.hourly_rate)])
        set_cell_lines(row.cells[3], [format_currency(session.total_pay)])

    for row, label, amount in zip(
        totals_table.rows,
        [
            "Subtotal",
            "Tax (0%)",
            "Invoice Total",
            "Less: Cash Payment Received",
            "BALANCE DUE (USD)",
        ],
        [
            format_currency(total),
            "$0.00",
            format_currency(total),
            "-$0.00",
            format_currency(total),
        ],
    ):
        bold_lines = {0} if label.startswith("BALANCE") else set()
        set_cell_lines(row.cells[0], [label], bold_lines=bold_lines)
        set_cell_lines(row.cells[1], [amount], bold_lines=bold_lines)

    for paragraph in document.paragraphs:
        if paragraph.text == "Paid in full by cash. Thank you.":
            set_paragraph_text(
                paragraph, f"Payment due by {due_date:%B %d, %Y}. Thank you."
            )
        elif paragraph.text.startswith("Website hosting and servicing"):
            set_paragraph_text(paragraph, "")

    for section in document.sections:
        footer_paragraph = section.footer.paragraphs[0]
        set_paragraph_text(
            footer_paragraph,
            "Thank you for your business.\n"
            "Please remit payment through PayPal to jameshardy1@gmail.com",
        )

    document.save(str(output_path))
    return total


def format_currency(amount: Decimal) -> str:
    return f"${amount:.2f}"


def set_cell_lines(
    cell: _Cell, lines: list[str], bold_lines: set[int] | None = None
) -> None:
    paragraphs = cell.paragraphs
    while len(paragraphs) < len(lines):
        cell.add_paragraph()
        paragraphs = cell.paragraphs
    for index, paragraph in enumerate(paragraphs):
        text = lines[index] if index < len(lines) else ""
        set_paragraph_text(
            paragraph, text, bold=bool(bold_lines and index in bold_lines)
        )


def set_paragraph_text(
    paragraph: Paragraph, text: str, bold: bool = False
) -> None:
    original_properties = (
        deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs else None
    )
    paragraph.clear()
    run = paragraph.add_run(text)
    if original_properties is not None:
        run._r.insert(0, original_properties)
    run.bold = bold
